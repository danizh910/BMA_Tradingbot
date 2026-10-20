"""Exportiert Arbeitsjournal und Prompt-Sammlung als Excel im Format der Schulvorlage.

Aufruf:  uv run python scripts/export_prozessunterlagen.py [--von 2026-10-01] [--bis 2026-10-28] [--out pfad.xlsx]

Quelle: docs/arbeitsjournal.csv und docs/prompt_sammlung.csv (Datum ISO JJJJ-MM-TT).
Ziel:   results/prozess/arbeitsjournal[_von_bis].xlsx, Blätter «Arbeitsjournal» und «Prompt-Log».
WARUM: Die CSV-Dateien sind die Hauptversion (gut lesbar im Git-Verlauf). Für die Abgaben
(AAR-Termine, Prozessunterlagen) braucht es dieselben Spalten wie die Schulvorlage
Arbeitsjournal_V1_VorlageBMA2026.xlsx: Querformat, Zeilenumbruch, Datum TT.MM.JJJJ.
"""

import argparse
import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from bot.config import ROOT

# CSV-Spalte -> (Titel, Erklärung, Spaltenbreite) wie in der Schulvorlage
JOURNAL_COLUMNS = {
    "datum": ("WANN?", "Datum", 11),
    "wer": ("WER?", "Wer hat etwas getan? \nAlle oder Einzelperson", 16.75),
    "was_warum_wie": ("WAS WARUM WIE? (EIGENLEISTUNG)",
                      "Was wurde warum wie getan? Probleme und Lösungsansätze.", 65),
    "woher_quellen_ki_dritte": ("WOHER? (QUELLEN, HILFE VON KI, DRITTPERSONEN)",
                                "Woher stammen die Informationen oder Materialien? Wie wurden Beiträge "
                                "von KI oder Drittpersonen generiert und eingebunden?", 45),
    "miss_erfolge": ("(MISS-)ERFOLGE",
                     "Was wurde erledigt, ging wie geplant oder besser? Was klappte nicht, wie geplant? "
                     "Woran lag es?", 40),
    "lehren_massnahmen": ("LEHREN/MASSNAHMEN",
                          "Was für Lehren ziehen Sie daraus? Welche Massnahmen ergreifen Sie, damit Sie "
                          "Erfolge in Zukunft wiederholen und Misserfolge vermeiden können?", 40),
}
PROMPT_COLUMNS = {
    "nr": ("Nr.", None, 6),
    "datum": ("Datum", None, 11),
    "tool_modell": ("Tool / Modell", None, 22),
    "wer": ("Wer", None, 16),
    "zweck": ("Zweck", None, 22),
    "prompt_wortlaut_oder_verweis": ("Prompt (Wortlaut oder Kurzfassung)", None, 70),
    "ergebnis_verwendung": ("Ergebnis / Verwendung", None, 40),
    "uebernommen": ("Übernommen?", None, 14),
    "iterationen": ("Iterationen", None, 11),
}


def read_csv(path: Path) -> list[dict]:
    """Liest eine Doku-CSV als Liste von Zeilen (dict pro Zeile)."""
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def filter_by_date(rows: list[dict], von: date | None, bis: date | None) -> list[dict]:
    """Behält nur Zeilen mit Datum in [von, bis].

    WARUM: Für einen AAR-Termin zählt nur die Arbeit seit dem letzten Termin.
    Zeilen ohne Datum fallen weg, sobald ein Filter gesetzt ist (Zeitpunkt unbekannt).
    """
    if von is None and bis is None:
        return rows
    kept = []
    for row in rows:
        if not row["datum"]:
            continue
        d = date.fromisoformat(row["datum"])
        if (von is None or d >= von) and (bis is None or d <= bis):
            kept.append(row)
    return kept


def write_sheet(ws, columns: dict, rows: list[dict]) -> None:
    """Schreibt Kopf (Titel + evtl. Erklärungszeile) und Zeilen im Layout der Vorlage.

    Datum wird als echtes Excel-Datum mit Format TT.MM.JJJJ geschrieben, damit Excel sortieren kann.
    """
    wrap = Alignment(wrap_text=True, vertical="top")
    ws.append([title for title, _, _ in columns.values()])
    header_rows = 1
    if any(note for _, note, _ in columns.values()):
        ws.append([note or "" for _, note, _ in columns.values()])
        header_rows = 2
    for row in rows:
        values = []
        for key in columns:
            value = row.get(key, "")
            if key == "datum" and value:
                value = date.fromisoformat(value)
            values.append(value)
        ws.append(values)

    for i, (_, _, width) in enumerate(columns.values(), start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    date_col = list(columns).index("datum") + 1
    for r, cells in enumerate(ws.iter_rows(), start=1):
        for cell in cells:
            cell.alignment = wrap
            if r <= header_rows:
                cell.font = Font(bold=(r == 1))
            elif cell.column == date_col:
                cell.number_format = "DD.MM.YYYY"

    ws.freeze_panes = ws.cell(header_rows + 1, 1)
    ws.print_title_rows = f"1:{header_rows}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def export(journal_csv: Path, prompts_csv: Path, out: Path,
           von: date | None = None, bis: date | None = None) -> Path:
    """Erstellt die Excel-Datei mit den Blättern «Arbeitsjournal» und «Prompt-Log» und gibt den Pfad zurück."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Arbeitsjournal"
    write_sheet(ws, JOURNAL_COLUMNS, filter_by_date(read_csv(journal_csv), von, bis))
    write_sheet(wb.create_sheet("Prompt-Log"), PROMPT_COLUMNS, filter_by_date(read_csv(prompts_csv), von, bis))
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main() -> None:
    """Liest die Optionen, exportiert und meldet den Pfad der Excel-Datei."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--von", type=date.fromisoformat, help="erstes Datum (JJJJ-MM-TT), inklusive")
    parser.add_argument("--bis", type=date.fromisoformat, help="letztes Datum (JJJJ-MM-TT), inklusive")
    parser.add_argument("--out", type=Path, help="Zieldatei (Standard: results/prozess/…)")
    args = parser.parse_args()

    suffix = f"_{args.von or ''}_{args.bis or ''}" if (args.von or args.bis) else ""
    out = args.out or ROOT / "results" / "prozess" / f"arbeitsjournal{suffix}.xlsx"
    path = export(ROOT / "docs" / "arbeitsjournal.csv", ROOT / "docs" / "prompt_sammlung.csv", out, args.von, args.bis)
    print(f"Exportiert: {path}")


if __name__ == "__main__":
    main()
