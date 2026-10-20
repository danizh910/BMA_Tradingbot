"""Tests des Excel-Exports der Prozessunterlagen (Format Schulvorlage, Datumsfilter)."""

import csv
import importlib.util
from datetime import date, datetime

import openpyxl

from bot.config import ROOT

spec = importlib.util.spec_from_file_location("export_prozessunterlagen", ROOT / "scripts" / "export_prozessunterlagen.py")
exp = importlib.util.module_from_spec(spec)
spec.loader.exec_module(exp)


def write_csv(path, columns, rows):
    """Schreibt eine kleine Test-CSV mit den gegebenen Spalten."""
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=columns)
        w.writeheader()
        w.writerows(rows)


def test_export_layout_and_date_filter(tmp_path):
    """Spalten wie Vorlage, Datum TT.MM.JJJJ, Querformat, Umbruch; Filter von/bis inklusive."""
    journal = tmp_path / "j.csv"
    prompts = tmp_path / "p.csv"
    write_csv(journal, list(exp.JOURNAL_COLUMNS), [
        {"datum": d, "wer": "Team", "was_warum_wie": f"Arbeit {d}", "woher_quellen_ki_dritte": "KI",
         "miss_erfolge": "ok", "lehren_massnahmen": "weiter"}
        for d in ["2026-09-30", "2026-10-19", "2026-10-28", "2026-11-01"]
    ])
    write_csv(prompts, list(exp.PROMPT_COLUMNS), [
        {"nr": "1", "datum": "2026-10-20", "zweck": "Doku"},
        {"nr": "2", "datum": "", "zweck": "ohne Datum"},
    ])

    out = exp.export(journal, prompts, tmp_path / "out" / "a.xlsx", date(2026, 10, 1), date(2026, 10, 28))
    wb = openpyxl.load_workbook(out)
    ws = wb["Arbeitsjournal"]

    assert [c.value for c in ws[1]] == ["WANN?", "WER?", "WAS WARUM WIE? (EIGENLEISTUNG)",
                                        "WOHER? (QUELLEN, HILFE VON KI, DRITTPERSONEN)",
                                        "(MISS-)ERFOLGE", "LEHREN/MASSNAHMEN"]
    data = list(ws.iter_rows(min_row=3, values_only=True))
    assert [r[0] for r in data] == [datetime(2026, 10, 19), datetime(2026, 10, 28)]
    assert ws["A3"].number_format == "DD.MM.YYYY"
    assert ws["C3"].alignment.wrap_text
    assert ws.page_setup.orientation == "landscape"

    p = wb["Prompt-Log"]
    assert p["A1"].value == "Nr."
    assert [r[4] for r in p.iter_rows(min_row=2, values_only=True)] == ["Doku"]


def test_no_filter_keeps_all_rows():
    """Ohne von/bis bleiben alle Zeilen, auch solche ohne Datum."""
    rows = [{"datum": ""}, {"datum": "2026-01-01"}]
    assert exp.filter_by_date(rows, None, None) == rows


def test_repo_csv_columns_match_export():
    """Die echten Doku-CSVs haben genau die Spalten, die der Export erwartet."""
    for name, columns in [("arbeitsjournal.csv", exp.JOURNAL_COLUMNS), ("prompt_sammlung.csv", exp.PROMPT_COLUMNS)]:
        with open(ROOT / "docs" / name, encoding="utf-8", newline="") as f:
            assert next(csv.reader(f)) == list(columns)
